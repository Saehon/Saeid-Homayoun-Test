from __future__ import annotations
import argparse, hashlib, importlib.util, io, json, math, zipfile
from pathlib import Path
import numpy as np
import pandas as pd
import requests
import statsmodels.api as sm
from scipy.stats import spearmanr
from statsmodels.stats.multitest import multipletests
from docx import Document

HERE=Path(__file__).resolve().parent
PARENT=HERE.parent
spec=importlib.util.spec_from_file_location("mnsc_base",PARENT/"run_study.py")
base=importlib.util.module_from_spec(spec); spec.loader.exec_module(base)
UA="MNSc-FIZ-CIZ-V3 academic replication"
DAM={
 "damodaran_betas_2026.xls":"https://www.stern.nyu.edu/~adamodar/pc/datasets/betas.xls",
 "damodaran_wacc_2026.xls":"https://www.stern.nyu.edu/~adamodar/pc/datasets/wacc.xls",
 "damodaran_capex_2026.xls":"https://www.stern.nyu.edu/~adamodar/pc/datasets/capex.xls",
 "damodaran_indreg_2026.xls":"https://www.stern.nyu.edu/~adamodar/pc/datasets/indreg.xls"}
SAFE={"Mkt-RF":"mktrf","SMB":"smb","HML":"hml","RMW":"rmw","CMA":"cma","RF":"rf",
      "SMALL LoBM":"sl","ME1 BM2":"sm","SMALL HiBM":"sh","BIG LoBM":"bl","ME2 BM2":"bm","BIG HiBM":"bh"}

def dl(url,path):
 r=requests.get(url,timeout=120,headers={"User-Agent":UA}); r.raise_for_status(); path.write_bytes(r.content)
 return {"file":path.name,"url":url,"bytes":len(r.content),"sha256":hashlib.sha256(r.content).hexdigest(),"source_class":"official-primary"}

def align():
 off,op,os=base.load_snapshot(2024); nf,np_,ns=base.load_snapshot(2025)
 off,nf=base._align(off,nf); op,np_=base._align(op,np_)
 common=sorted(set(off.date)&set(nf.date)&set(op.date)&set(np_.date))
 ff=pd.DataFrame({"date":common})
 for tag,f,p in [("old",off,op),("new",nf,np_)]:
  f=f.set_index("date").loc[common]; p=p.set_index("date").loc[common]
  for c in base.FACTORS: ff[f"{tag}_{SAFE[c]}"]=f[c].to_numpy()
  for c in base.PORTS: ff[f"{tag}_{SAFE[c]}"]=p[c].to_numpy()
 for v in SAFE.values(): ff[f"d_{v}"]=ff[f"new_{v}"]-ff[f"old_{v}"]
 ff["mdate"]=ff.date.dt.year*12+ff.date.dt.month-1
 return ff,os+ns

def table1(ff):
 out=[]
 for label,v in SAFE.items():
  a,b=ff[f"old_{v}"],ff[f"new_{v}"]; d=b-a
  out.append(dict(series=label,code=v,n=len(d),old_mean_pct=100*a.mean(),new_mean_pct=100*b.mean(),
   old_sd_pct=100*a.std(ddof=1),new_sd_pct=100*b.std(ddof=1),mean_diff_bps=10000*d.mean(),
   DCS_mean_abs_bps=10000*d.abs().mean(),max_abs_diff_bps=10000*d.abs().max(),old_new_corr=a.corr(b)))
 return pd.DataFrame(out)

def table2(ff):
 vs=[f"d_{v}" for v in SAFE.values()]; out=[]
 for i,a in enumerate(vs):
  for b in vs[i:]:
   out.append(dict(var1=a,var2=b,pearson=ff[a].corr(ff[b]),spearman=spearmanr(ff[a],ff[b],nan_policy="omit").statistic,n=len(ff)))
 return pd.DataFrame(out)

def table3(ff):
 models={"CAPM":["mktrf"],"FF3":["mktrf","smb","hml"],"FF5":["mktrf","smb","hml","rmw","cma"]}; out=[]
 for model,xs in models.items():
  for p in [SAFE[x] for x in base.PORTS]:
   res={}
   for tag in ["old","new"]:
    y=ff[f"{tag}_{p}"]-ff[f"{tag}_rf"]; X=sm.add_constant(pd.DataFrame({x:ff[f"{tag}_{x}"] for x in xs}),has_constant="add")
    m=sm.OLS(y,X).fit(cov_type="HAC",cov_kwds={"maxlags":6})
    res[tag]=dict(a=m.params["const"],t=m.tvalues["const"],pv=m.pvalues["const"],r2=m.rsquared_adj,n=m.nobs)
   o,n=res["old"],res["new"]
   out.append(dict(model=model,portfolio=p,old_alpha_month=o["a"],new_alpha_month=n["a"],old_alpha_annual=12*o["a"],new_alpha_annual=12*n["a"],
    old_t=o["t"],new_t=n["t"],old_p=o["pv"],new_p=n["pv"],old_adj_r2=o["r2"],new_adj_r2=n["r2"],n=int(o["n"]),
    alpha_change_bps=10000*(n["a"]-o["a"]),sign_change=np.sign(n["a"])!=np.sign(o["a"]),
    significance_change_5pct=(n["pv"]<.05)!=(o["pv"]<.05),conclusion_reversal=(np.sign(n["a"])!=np.sign(o["a"])) or ((n["pv"]<.05)!=(o["pv"]<.05))))
 return pd.DataFrame(out)

def xls(path):
 p=pd.read_excel(path,header=None); hi=None
 for i in range(min(25,len(p))):
  if "industry name" in " ".join(str(x).lower() for x in p.iloc[i] if pd.notna(x)): hi=i; break
 if hi is None: raise ValueError(f"Industry header not found: {path}")
 d=pd.read_excel(path,header=hi); d.columns=[str(c).strip() for c in d.columns]; ic=next(c for c in d if "industry name" in c.lower()); return d.rename(columns={ic:"industry"})
def col(d,*keys):
 for c in d:
  s=c.lower().replace("\n"," ")
  if all(k in s for k in keys): return c
 raise KeyError((keys,list(d.columns)))
def num(s): return pd.to_numeric(s.astype(str).str.replace("%","",regex=False).str.replace(",","",regex=False),errors="coerce")
def damodaran(raw):
 man=[]; z={}
 for fn,u in DAM.items(): man.append(dl(u,raw/fn)); z[fn]=xls(raw/fn)
 b,w,c,a=z["damodaran_betas_2026.xls"],z["damodaran_wacc_2026.xls"],z["damodaran_capex_2026.xls"],z["damodaran_indreg_2026.xls"]
 bc=col(b,"beta"); nc=next((x for x in b if "number of firms" in x.lower()),None)
 if nc is None: nc=col(w,"number of firms"); bb=b[["industry",bc]].copy(); nn=w[["industry",nc]].copy(); nn.columns=["industry","n_firms"]; bb.columns=["industry","beta"]; bb=bb.merge(nn,on="industry")
 else: bb=b[["industry",bc,nc]].copy(); bb.columns=["industry","beta","n_firms"]
 wc=col(w,"cost of capital"); cc=next(x for x in c if "net cap" in x.lower() and "sales" in x.lower()); ac=next(x for x in a if "jensen" in x.lower() and "alpha" in x.lower())
 ww=w[["industry",wc]].copy(); ww.columns=["industry","wacc"]; ccx=c[["industry",cc]].copy(); ccx.columns=["industry","netcapex_sales"]; aa=a[["industry",ac]].copy(); aa.columns=["industry","jensen_alpha"]
 d=bb.merge(ww,on="industry").merge(ccx,on="industry").merge(aa,on="industry")
 for q in ["beta","n_firms","wacc","netcapex_sales","jensen_alpha"]: d[q]=num(d[q])
 for q in ["wacc","netcapex_sales","jensen_alpha"]:
  if d[q].abs().median()>1: d[q]/=100
 d["ln_n_firms"]=np.log(d["n_firms"].where(d["n_firms"]>0)); d=d.dropna(subset=["jensen_alpha","beta","wacc","netcapex_sales","ln_n_firms"])
 return d,man

def table4(d):
 specs={"M1":["beta"],"M2":["beta","wacc"],"M3":["beta","wacc","netcapex_sales"],"M4":["beta","wacc","netcapex_sales","ln_n_firms"]}; out=[]
 for name,xs in specs.items():
  m=sm.OLS(d.jensen_alpha,sm.add_constant(d[xs],has_constant="add")).fit(cov_type="HC3")
  for q in ["const"]+xs: out.append(dict(model=name,term=q,coef=m.params[q],se_hc3=m.bse[q],t=m.tvalues[q],p=m.pvalues[q],n=int(m.nobs),r2=m.rsquared,adj_r2=m.rsquared_adj))
 return pd.DataFrame(out)
def hmean(s,L):
 m=sm.OLS(s,np.ones((len(s),1))).fit(cov_type="HAC",cov_kwds={"maxlags":L}); return m.params[0],m.bse[0],m.tvalues[0],m.pvalues[0]
def table5(ff):
 out=[]; vs=[f"d_{v}" for v in SAFE.values()]; idx=[]; ps=[]
 for L in [3,6,12]:
  for v in vs:
   b,se,t,p=hmean(ff[v],L); out.append(dict(panel="HAC_mean_revision",series=v,spec=f"NW_lag_{L}",estimate_bps=10000*b,se_bps=10000*se,t=t,p=p,q_bh=np.nan,n=len(ff)))
   if L==6: idx.append(len(out)-1); ps.append(p)
 q=multipletests(ps,method="fdr_bh")[1]
 for i,x in zip(idx,q): out[i]["q_bh"]=x
 periods={"full":pd.Series(True,index=ff.index),"post_1990":ff.date>=pd.Timestamp("1990-01-01"),"post_2000":ff.date>=pd.Timestamp("2000-01-01"),"post_gfc":ff.date>=pd.Timestamp("2009-01-01"),"post_covid":ff.date>=pd.Timestamp("2020-01-01")}
 for name,mask in periods.items():
  for v in vs: out.append(dict(panel="subperiod_DCS",series=v,spec=name,estimate_bps=10000*ff.loc[mask,v].abs().mean(),se_bps=np.nan,t=np.nan,p=np.nan,q_bh=np.nan,n=int(mask.sum())))
 return pd.DataFrame(out)
def stata(d,path):
 z=d.copy()
 if "date" in z: z["date_str"]=z.date.dt.strftime("%Y-%m-%d"); z=z.drop(columns="date")
 z.to_stata(path,write_index=False,version=118)
def docx(tables,path,ff):
 d=Document(); d.add_heading("Management Science Empirical Package — Five Canonical Tables",0); d.add_paragraph(f"Matched FIZ/CIZ sample: {ff.date.min().date()}–{ff.date.max().date()} ({len(ff)} months). Damodaran January-2026 industry data are external validation only.")
 titles=["Table 1. Descriptive Statistics and DCS","Table 2. Correlations of Revisions","Table 3. CAPM/FF3/FF5 HAC Regressions","Table 4. Damodaran External Validation","Table 5. Robustness, FDR, and Subperiod DCS"]
 for title,z in zip(titles,tables):
  d.add_heading(title,1); t=d.add_table(rows=1,cols=len(z.columns)); t.style="Table Grid"
  for j,c in enumerate(z): t.rows[0].cells[j].text=str(c)
  for _,r in z.iterrows():
   cells=t.add_row().cells
   for j,c in enumerate(z):
    x=r[c]; cells[j].text=("" if pd.isna(x) else f"{x:.4f}") if isinstance(x,(float,np.floating)) else str(x)
 d.save(path)
def main():
 ap=argparse.ArgumentParser(); ap.add_argument("--output",default="artifact"); A=ap.parse_args(); out=Path(A.output)
 raw=out/"01_Data_Raw"; clean=out/"02_Data_Clean"; code=out/"03_Stata_Code"; tabs=out/"04_Output_5_Tables"; ms=out/"05_Manuscript"; lit=out/"07_FT50_ABS4_Literature"
 for p in [raw,clean,code,tabs,ms,lit]: p.mkdir(parents=True,exist_ok=True)
 ff,src=align(); d,ds=damodaran(raw)
 for y in [2024,2025]:
  for fn,fam in [(base.FF5_FILE,"ff5"),(base.PORT6_FILE,"port6")]: src.append(dl(base.BASE.format(year=y,file=fn),raw/f"fama_{y}_{fam}.zip"))
 pd.DataFrame(src+ds).drop_duplicates(subset=["url"]).to_csv(raw/"source_manifest.csv",index=False)
 ff.to_csv(clean/"fiz_ciz_monthly.csv",index=False); stata(ff,clean/"fiz_ciz_monthly.dta"); d.to_csv(clean/"damodaran_industry_2026.csv",index=False); stata(d,clean/"damodaran_industry_2026.dta")
 T=[table1(ff),table2(ff),table3(ff),table4(d),table5(ff)]; names=["table1_descriptive_dcs.csv","table2_correlations.csv","table3_baseline_regressions.csv","table4_damodaran_external_validation.csv","table5_robustness_fdr_subperiods.csv"]
 for z,n in zip(T,names): z.to_csv(tabs/n,index=False)
 docx(T,ms/"Empirical_Results_5_Tables.docx",ff)
 import shutil; shutil.copy(HERE/"stata"/"00_master.do",code/"00_master.do"); shutil.copy(HERE/"STATA_CODEX_AGENT.md",code/"STATA_CODEX_AGENT.md")
 refs=pd.DataFrame([
 ["He, Huang, Li & Zhou","Shrinking Factor Dimension: A Reduced-Rank Approach","Management Science","10.1287/mnsc.2022.4563"],
 ["Chen","Do t-Statistic Hurdles Need to Be Raised?","Management Science","10.1287/mnsc.2023.03083"],
 ["Feng, Giglio & Xiu","Taming the Factor Zoo","Journal of Finance","10.1111/jofi.12883"],
 ["Giglio, Xiu & Zhang","Test Assets and Weak Factors","Journal of Finance","10.1111/jofi.13415"],
 ["Gu, Kelly & Xiu","Empirical Asset Pricing via Machine Learning","Review of Financial Studies","10.1093/rfs/hhaa009"],
 ["Hou, Xue & Zhang","Replicating Anomalies","Review of Financial Studies","10.1093/rfs/hhy131"]],columns=["authors","title","journal","doi"]); refs.to_csv(lit/"FT50_ABS4_method_anchors.csv",index=False)
 (out/"README.md").write_text(f"# MNSc FIZ-CIZ V3 real-data package\n\nMatched FIZ/CIZ sample: {ff.date.min().date()}–{ff.date.max().date()} ({len(ff)} months).\nDamodaran complete-case industries: {len(d)}.\n\nFive canonical tables are in `04_Output_5_Tables`. Raw official files and SHA-256 provenance are in `01_Data_Raw`.\n",encoding="utf-8")
 (out/"RUN_SUMMARY.json").write_text(json.dumps({"common_start":str(ff.date.min().date()),"common_end":str(ff.date.max().date()),"months":len(ff),"damodaran_industries":len(d),"table_rows":[len(x) for x in T]},indent=2),encoding="utf-8")
 print((out/"RUN_SUMMARY.json").read_text())
if __name__=="__main__": main()
